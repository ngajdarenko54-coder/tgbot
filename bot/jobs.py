"""Загрузка базы вакансий из Excel и подбор вакансий по ответам анкеты."""
from dataclasses import dataclass, field
from urllib.parse import urlencode, urlparse, urlunparse, parse_qsl

from openpyxl import load_workbook

REQUIRED_COLUMNS = [
    "offer_id", "title", "category", "age_min", "need_license_b", "need_license_c",
    "need_car", "hours_min", "schedule", "remote", "districts", "description",
    "pay_text", "requirements", "link", "priority", "active",
]


def _split(value) -> frozenset[str]:
    if value is None:
        return frozenset()
    return frozenset(x.strip().lower() for x in str(value).split(",") if x.strip())


def _int(value, default: int = 0) -> int:
    if value is None or str(value).strip() == "":
        return default
    return int(float(value))


@dataclass(frozen=True)
class Job:
    offer_id: int
    title: str
    categories: frozenset[str]
    age_min: int
    need_license_b: bool
    need_license_c: bool
    need_car: bool
    hours_min: int
    schedules: frozenset[str]
    remote: bool
    districts: frozenset[str]
    description: str
    pay_text: str
    requirements: str
    link: str
    priority: int
    active: bool
    advertiser: str = ""

    @property
    def erid(self) -> str:
        """Токен маркировки рекламы из партнёрской ссылки (?erid=...)."""
        return dict(parse_qsl(urlparse(self.link).query)).get("erid", "")

    def link_with_subid(self, param: str, subid: str) -> str:
        parts = urlparse(self.link)
        query = dict(parse_qsl(parts.query))
        query[param] = subid
        return urlunparse(parts._replace(query=urlencode(query)))


def load_jobs(path: str) -> list[Job]:
    """Читает лист jobs. Бросает ValueError с понятным текстом, если файл кривой."""
    wb = load_workbook(path, read_only=True, data_only=True)
    if "jobs" not in wb.sheetnames:
        raise ValueError("В файле нет листа «jobs»")
    rows = list(wb["jobs"].iter_rows(values_only=True))
    if not rows:
        raise ValueError("Лист «jobs» пустой")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ValueError("Не хватает колонок: " + ", ".join(missing))
    idx = {name: header.index(name) for name in REQUIRED_COLUMNS}
    adv_idx = header.index("advertiser") if "advertiser" in header else None

    jobs: list[Job] = []
    for n, row in enumerate(rows[1:], start=2):
        if row is None or all(v is None for v in row):
            continue
        get = lambda col: row[idx[col]] if idx[col] < len(row) else None  # noqa: E731
        try:
            jobs.append(Job(
                offer_id=_int(get("offer_id")),
                title=str(get("title") or "").strip(),
                categories=_split(get("category")),
                age_min=_int(get("age_min"), 18),
                need_license_b=bool(_int(get("need_license_b"))),
                need_license_c=bool(_int(get("need_license_c"))),
                need_car=bool(_int(get("need_car"))),
                hours_min=_int(get("hours_min")),
                schedules=_split(get("schedule")),
                remote=bool(_int(get("remote"))),
                districts=_split(get("districts")),
                description=str(get("description") or "").strip(),
                pay_text=str(get("pay_text") or "").strip(),
                requirements=str(get("requirements") or "").strip(),
                link=str(get("link") or "").strip(),
                priority=_int(get("priority"), 999),
                active=bool(_int(get("active"), 1)),
                advertiser=str(row[adv_idx] or "").strip() if adv_idx is not None and adv_idx < len(row) else "",
            ))
        except (TypeError, ValueError) as e:
            raise ValueError(f"Строка {n}: не получилось прочитать ({e})") from e
    if not jobs:
        raise ValueError("В файле нет ни одной вакансии")
    return jobs


@dataclass
class Answers:
    """Ответы анкеты. None / 'any' = неважно."""
    age: int = 18
    category: str = "any"
    license_b: bool = False
    license_c: bool = False
    has_car: bool = False
    hours: int = 20          # сколько часов в неделю готов работать
    schedule: str = "any"    # flex | 2/2 | 5/2 | night | full | vahta | any
    remote: str = "any"      # only | no | any
    district: str = "any"
    relaxed: list[str] = field(default_factory=list)


def _fits(job: Job, a: Answers, skip: set[str]) -> bool:
    if not job.active or not job.link:
        return False
    if a.age < job.age_min:
        return False
    if job.need_license_b and not a.license_b:
        return False
    if job.need_license_c and not a.license_c:
        return False
    if job.need_car and not a.has_car:
        return False
    if a.category != "any" and a.category not in job.categories:
        return False
    if a.remote == "only" and not job.remote:
        return False
    if a.remote == "no" and job.categories <= {"remote"}:
        return False
    if "hours" not in skip and a.hours < job.hours_min:
        return False
    if "schedule" not in skip and a.schedule != "any" and job.schedules and a.schedule not in job.schedules:
        return False
    if "district" not in skip and a.district != "any" and job.districts and a.district.lower() not in job.districts:
        return False
    return True


# Порядок ослабления фильтров, если ничего не нашлось. Жёсткие требования
# (возраст, права, авто, категория) не ослабляются никогда.
RELAX_ORDER = ["district", "schedule", "hours"]


def match_jobs(jobs: list[Job], a: Answers, limit: int = 3) -> list[Job]:
    skip: set[str] = set()
    a.relaxed = []
    for step in [None, *RELAX_ORDER]:
        if step:
            skip.add(step)
            a.relaxed.append(step)
        found = sorted((j for j in jobs if _fits(j, a, skip)), key=lambda j: j.priority)
        if found:
            return found[:limit]
    return []
