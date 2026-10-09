"""Тесты без доступа к Telegram: подбор вакансий и полный диалог на подменённой сессии."""
import asyncio
import shutil
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from aiogram import Bot, Dispatcher
from aiogram.client.default import DefaultBotProperties
from aiogram.client.session.base import BaseSession
from aiogram.methods import EditMessageText, SendMessage, TelegramMethod
from aiogram.types import CallbackQuery, Chat, Message, Update, User

from bot.config import Config
from bot.handlers import JobsHolder, router, send_due_followups
from bot.jobs import Answers, load_jobs, match_jobs
from bot.storage import Storage

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture()
def xlsx(tmp_path):
    """Копия базы, где ссылки есть у всех вакансий (в реальном файле у 2465 её нет)."""
    path = tmp_path / "jobs.xlsx"
    shutil.copy(ROOT / "data" / "jobs.xlsx", path)
    wb = load_workbook(path)
    ws = wb["jobs"]
    col = [c.value for c in ws[1]].index("link") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(r, col, f"https://example.com/offer/{ws.cell(r, 1).value}?a=1")
    wb.save(path)
    return str(path)


def test_real_file_links():
    jobs = {j.offer_id: j for j in load_jobs(str(ROOT / "data" / "jobs.xlsx"))}
    assert len(jobs) == 16
    with_links = {i for i, j in jobs.items() if j.link}
    assert len(with_links) == 13 and 2465 not in with_links
    for i in with_links:
        assert jobs[i].link.startswith("https://trk.ppdu.ru/click/") and jobs[i].erid
    shown = match_jobs(list(jobs.values()), Answers(age=20), limit=50)
    ids = {j.offer_id for j in shown}
    assert 2465 not in ids and 2682 not in ids  # без ссылки / выключен
    courier = match_jobs(list(jobs.values()), Answers(age=19, category="courier", hours=20, schedule="flex"))
    assert courier[0].offer_id == 2304


def test_subid_keeps_erid_and_card_has_ad_label():
    from bot.handlers import job_card
    job = next(j for j in load_jobs(str(ROOT / "data" / "jobs.xlsx")) if j.offer_id == 1570)
    url = job.link_with_subid("sub1", "abc123")
    assert "erid=2SDnjcXP37s" in url and url.endswith("sub1=abc123")
    assert "Реклама" in job_card(job) and "erid: 2SDnjcXP37s" in job_card(job)


def test_courier_without_license_gets_yandex_first(xlsx):
    jobs = load_jobs(xlsx)
    found = match_jobs(jobs, Answers(age=19, category="courier", hours=20, schedule="flex"))
    assert found[0].offer_id == 2304
    assert all("courier" in j.categories for j in found)
    assert 2682 not in [j.offer_id for j in found]  # Купер выключен (active=0)


def test_truck_driver_needs_license_c(xlsx):
    jobs = load_jobs(xlsx)
    no_c = match_jobs(jobs, Answers(age=25, category="driver", license_b=True, hours=60, schedule="full"))
    with_c = match_jobs(jobs, Answers(age=25, category="driver", license_b=True, license_c=True,
                                      hours=60, schedule="full"))
    assert no_c == [] and [j.offer_id for j in with_c] == [2864]


def test_relaxes_schedule_but_not_category(xlsx):
    jobs = load_jobs(xlsx)
    a = Answers(age=20, category="remote", hours=10, schedule="night")
    found = match_jobs(jobs, a)
    assert found and "schedule" in a.relaxed
    assert all("remote" in j.categories for j in found)


def test_only_remote(xlsx):
    jobs = load_jobs(xlsx)
    found = match_jobs(jobs, Answers(age=20, remote="only", hours=10), limit=20)
    assert found and all(j.remote for j in found)


def test_subid_added_to_link(xlsx):
    job = load_jobs(xlsx)[0]
    assert job.link_with_subid("sub1", "abc") == f"https://example.com/offer/{job.offer_id}?a=1&sub1=abc"


def test_bad_file_rejected(tmp_path):
    from openpyxl import Workbook
    wb = Workbook(); wb.active.title = "jobs"; wb.active.append(["offer_id", "title"])
    p = tmp_path / "bad.xlsx"; wb.save(p)
    with pytest.raises(ValueError, match="Не хватает колонок"):
        load_jobs(str(p))


# ---------- полный диалог на подменённой сессии ----------

class FakeSession(BaseSession):
    def __init__(self):
        super().__init__()
        self.sent: list[TelegramMethod] = []
        self.mid = 100

    async def make_request(self, bot, method, timeout=None):
        self.sent.append(method)
        if isinstance(method, (SendMessage, EditMessageText)):
            self.mid += 1
            chat_id = method.chat_id or 1
            return Message(message_id=self.mid, date=datetime.now(), text=method.text,
                           chat=Chat(id=chat_id, type="private"))
        return True

    async def close(self):
        pass

    async def stream_content(self, *a, **k):
        yield b""


USER = User(id=42, is_bot=False, first_name="Test")
CHAT = Chat(id=42, type="private")


def run_dialog(xlsx, tmp_path):
    session = FakeSession()
    bot = Bot("123:ABC", session=session, default=DefaultBotProperties(parse_mode="HTML"))
    cfg = Config("123:ABC", frozenset({42}), "https://t.me/channel", "sub1", 48,
                 str(tmp_path / "bot.db"), xlsx)
    storage, jobs = Storage(cfg.db_path), JobsHolder(xlsx)
    dp = Dispatcher()
    router._parent_router = None  # один router на несколько тестовых диспетчеров
    dp.include_router(router)
    dp["cfg"], dp["storage"], dp["jobs"] = cfg, storage, jobs
    uid = iter(range(1, 1000))

    async def text(t):
        msg = Message(message_id=next(uid), date=datetime.now(), chat=CHAT, from_user=USER, text=t)
        await dp.feed_update(bot, Update(update_id=next(uid), message=msg))

    async def click(data):
        msg = Message(message_id=1, date=datetime.now(), chat=CHAT, from_user=USER, text="q")
        cb = CallbackQuery(id=str(next(uid)), from_user=USER, chat_instance="x", message=msg, data=data)
        await dp.feed_update(bot, Update(update_id=next(uid), callback_query=cb))

    return session, bot, storage, jobs, text, click


def texts_of(session):
    return [m.text for m in session.sent if isinstance(m, (SendMessage, EditMessageText))]


def test_full_dialog(xlsx, tmp_path):
    session, bot, storage, jobs, text, click = run_dialog(xlsx, tmp_path)

    async def scenario():
        await text("/start")
        await click("age:18")
        await click("cat:courier")
        await click("lic:none")
        await click("hrs:20")
        await click("sch:flex")
        cards = texts_of(session)
        assert any("Курьер Яндекс.Еды" in t for t in cards)

        await click("apply:2304")
        apply_msg = [m for m in session.sent if isinstance(m, SendMessage) and m.reply_markup
                     and m.reply_markup.inline_keyboard[0][0].url]
        url = apply_msg[-1].reply_markup.inline_keyboard[0][0].url
        assert "sub1=" + storage.user_code(42) in url

        await click("remind:2304:1")
        storage.db.execute("UPDATE followups SET due_at=0")
        await send_due_followups(bot, storage, jobs)
        assert any("Как дела с вакансией" in t for t in texts_of(session))

        await click("fu:called:2304")
        assert storage.db.execute("SELECT stage FROM followups").fetchone()[0] == 2

        await text("/stats")
        assert any("Нажали «Откликнуться»: 1" in t for t in texts_of(session))

        await text("/delete")
        assert storage.db.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0

    asyncio.run(scenario())


def test_under_18_and_remote_question(xlsx, tmp_path):
    session, bot, storage, jobs, text, click = run_dialog(xlsx, tmp_path)

    async def scenario():
        await text("/start")
        await click("age:0")
        assert any("с 18 лет" in t for t in texts_of(session))
        await text("/start")
        for d in ("age:21", "cat:any", "lic:b", "car:0", "hrs:10", "sch:any"):
            await click(d)
        assert texts_of(session)[-1].startswith("Удалённая работа")
        await click("rem:only")
        assert any("Тетрик" in t for t in texts_of(session))

    asyncio.run(scenario())
